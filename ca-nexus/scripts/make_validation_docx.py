"""Convert docs/VALIDATION-PACK.md to an editable Word doc for the validator.
Usage: python scripts/make_validation_docx.py
Output: docs/VALIDATION-PACK.docx (checkboxes ☐/☒ editable as text).
"""
import os
import re

from docx import Document
from docx.shared import Pt

SRC = os.path.join(os.path.dirname(__file__), "..", "docs", "VALIDATION-PACK.md")
OUT = os.path.join(os.path.dirname(__file__), "..", "docs", "VALIDATION-PACK.docx")


def add_rich(par, text):
    """**bold** segments become bold runs; rest normal."""
    for i, seg in enumerate(re.split(r"(\*\*.+?\*\*)", text)):
        if not seg:
            continue
        run = par.add_run(seg[2:-2] if seg.startswith("**") and seg.endswith("**") else seg)
        if i % 2 == 1:
            run.bold = True


def main():
    lines = open(SRC, encoding="utf-8").read().splitlines()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(11)
    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.strip() == "---":
            doc.add_paragraph("─" * 40)
        elif re.match(r"^- \[[ x]\]", line.strip()):
            checked = line.strip().startswith("- [x]") or line.strip().startswith("- [X]")
            body = re.sub(r"^- \[[ xX]\]\s*", "", line.strip())
            par = doc.add_paragraph()
            box = par.add_run("☒ " if checked else "☐ ")
            box.font.size = Pt(13)
            add_rich(par, body)
        elif line.strip().startswith("- "):
            par = doc.add_paragraph(style="List Bullet")
            add_rich(par, line.strip()[2:])
        else:
            par = doc.add_paragraph()
            add_rich(par, line.strip())
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
